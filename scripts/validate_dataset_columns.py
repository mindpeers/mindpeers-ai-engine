#!/usr/bin/env python3
"""
Validate user feature dataset columns against Engine POC v1 contract.

Usage:
  .venv/bin/python scripts/validate_dataset_columns.py /path/to/your_dataset.xlsx
  .venv/bin/python scripts/validate_dataset_columns.py /path/to/your_dataset.csv

Optional: --sheet SheetName  (default: first sheet or User_Feature_Row_Sample)
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

REQUIRED_V1 = [
    "user_id",
    "as_of_date",
    "days_active",
    "system_type",
    "data_completeness_score",
    "feature_version",
    "mood_avg_14d",
    "mood_slope_7d",
    "mood_volatility_14d",
    "core_om_wellbeing",
    "core_om_problems",
    "core_om_functioning",
    "core_om_risk",
    "core_om_delta_30d",
    "gad7_normalized_latest",
    "motivation_avg_14d",
    "motivation_slope_7d",
    "confidence_avg_14d",
    "confidence_slope_7d",
    "engagement_rate_7d",
    "engagement_slope_7d",
    "engagement_delta_30d",
    "therapy_attendance_rate_30d",
    "sleep_avg_7d",
    "sleep_delta_30d",
    "sleep_slope_7d",
    "sleep_persistence_low_days",
    "resting_hr_relative",
    "activity_slope_7d",
    "cortisol_flag",
    "withdrawal_flag",
    "sleep_mood_coupled_decline",
]

ALIASES = {
    "avg_mood_14d": "mood_avg_14d",
}

RECOMMENDED = [
    "journaling_concern_score_7d",
    "composite_risk_percentile",
    "mood_delta_30d",
]

SNAPSHOT_COLUMNS = [
    "cognitive_readiness_score",
    "clarity_score",
    "emotional_balance_score",
    "resilience_score",
    "capacity_score",
    "trends_json",
]


def load_columns(path: Path, sheet: str | None) -> tuple[list[str], int, str | None]:
    if path.suffix.lower() == ".csv":
        import csv

        with path.open(newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            headers = next(reader)
            row_count = sum(1 for _ in reader)
        return [h.strip() for h in headers], row_count, None

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet:
        if sheet not in wb.sheetnames:
            wb.close()
            raise SystemExit(f"Sheet not found: {sheet}. Available: {wb.sheetnames}")
        ws = wb[sheet]
        sheet_name = sheet
    else:
        preferred = [
            "User_Feature_Row_Sample",
            "User_ML_Dataset_50K",
            "User_Column_Definitions",
        ]
        sheet_name = next((p for p in preferred if p in wb.sheetnames), wb.sheetnames[0])
        ws = wb[sheet_name]

    rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    row_count = max((ws.max_row or 1) - 1, 0)
    wb.close()
    return headers, row_count, sheet_name


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate dataset columns vs Engine POC v1")
    parser.add_argument("path", type=Path, help="Path to .xlsx or .csv")
    parser.add_argument("--sheet", default=None, help="Excel sheet name")
    args = parser.parse_args()

    if not args.path.exists():
        print(f"File not found: {args.path}", file=sys.stderr)
        return 1

    headers, row_count, sheet_name = load_columns(args.path, args.sheet)
    print(f"File: {args.path}")
    if sheet_name:
        print(f"Sheet: {sheet_name}")

    print(f"Rows (excl. header): {row_count:,}")
    print(f"Columns found: {len(headers)}")
    print()

    col_set = {h for h in headers if h}
    normalized = set(col_set)
    rename_notes = []
    for old, new in ALIASES.items():
        if old in col_set and new not in col_set:
            normalized.add(new)
            rename_notes.append(f"  Alias: `{old}` → treat as `{new}`")

    missing = [c for c in REQUIRED_V1 if c not in normalized]
    present_required = [c for c in REQUIRED_V1 if c in normalized]
    extra = sorted(col_set - set(REQUIRED_V1) - set(ALIASES.keys()) - set(RECOMMENDED))
    recommended_missing = [c for c in RECOMMENDED if c not in normalized]
    has_snapshot = [c for c in SNAPSHOT_COLUMNS if c in col_set]

    print("=== Required v1 L2 columns ===")
    print(f"Present: {len(present_required)} / {len(REQUIRED_V1)}")
    if missing:
        print("MISSING (add or compute):")
        for c in missing:
            print(f"  - {c}")
    else:
        print("All required L2 columns present (or aliased).")

    if rename_notes:
        print("\n=== Aliases detected ===")
        print("\n".join(rename_notes))

    if recommended_missing:
        print("\n=== Recommended (optional) missing ===")
        for c in recommended_missing:
            print(f"  - {c}")

    if has_snapshot:
        print("\n=== Note: output/score columns found on feature row ===")
        print("These belong on user_engine_snapshot after L3 batch, not required on L2 input:")
        for c in has_snapshot:
            print(f"  - {c}")

    if extra:
        print(f"\n=== Extra columns in your file ({len(extra)}) — OK if legacy/labels ===")
        for c in extra[:40]:
            print(f"  + {c}")
        if len(extra) > 40:
            print(f"  ... and {len(extra) - 40} more")

    print("\n=== Verdict ===")
    if missing:
        print("FAIL — missing required v1 columns. See Dataset-Column-Validation-Checklist.md")
        return 1
    print("PASS — column contract satisfied for L2 v1 Engine (proxy biological scope).")
    if row_count < 100:
        print(f"WARN — only {row_count} rows; you mentioned thousands — confirm correct sheet/file.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
